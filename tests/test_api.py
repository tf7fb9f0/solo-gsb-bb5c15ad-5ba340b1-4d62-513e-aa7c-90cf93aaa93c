"""HTTP API 端到端测试：上传 -> 查询影响面 -> 导出。"""
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app

client = TestClient(app)


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["A1"] = 3
    ws["B1"] = "=A1*2"
    ws2 = wb.create_sheet("S2")
    ws2["C3"] = "=S1!B1+1"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_health():
    assert client.get("/api/health").json()["status"] == "ok"


def test_index_page():
    r = client.get("/")
    assert r.status_code == 200 and "影响分析" in r.text


def test_upload_and_impact_flow():
    r = client.post(
        "/api/upload",
        files={"file": ("t.xlsx", make_xlsx(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200
    d = r.json()
    wid = d["workbook_id"]
    assert d["formula_count"] == 2
    assert any(s["name"] == "S1" for s in d["sheets"])

    r = client.get(f"/api/workbooks/{wid}/impact",
                   params={"cell": "S1!A1", "direction": "downstream"})
    assert r.status_code == 200
    imp = r.json()
    cells = {i["cell"] for i in imp["impacted"]}
    assert {"S1!B1", "S2!C3"} <= cells

    # 上游反向查询
    r = client.get(f"/api/workbooks/{wid}/impact",
                   params={"cell": "S2!C3", "direction": "upstream"})
    cells = {i["cell"] for i in r.json()["impacted"]}
    assert {"S1!B1", "S1!A1"} <= cells

    # 工作表筛选
    r = client.get(f"/api/workbooks/{wid}/cells", params={"sheet": "S2"})
    assert r.status_code == 200
    assert all(c["sheet"] == "S2" for c in r.json()["cells"])

    # 导出
    r = client.get(f"/api/workbooks/{wid}/export")
    assert r.status_code == 200
    export = r.json()
    assert export["filename"] == "t.xlsx"
    assert len(export["edges"]) == 2
    assert {"from": "S1!A1", "to": "S1!B1"} in export["edges"]
    assert "attachment" in r.headers["content-disposition"]


def test_bad_upload():
    r = client.post("/api/upload", files={"file": ("t.txt", b"hello", "text/plain")})
    assert r.status_code == 400


def test_missing_workbook():
    r = client.get("/api/workbooks/nonexistent")
    assert r.status_code == 404


def make_audit_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["A1"] = 1
    ws["A3"] = "x"                  # 扩展已用范围，使 B1 在范围内但为空
    ws["C1"] = "=A1+B1"             # B1 空引用
    ws["C2"] = "=幽灵!A1"           # 缺失工作表
    ws["C3"] = "=A999"              # 越界
    ws2 = wb.create_sheet("数据")
    for i in range(1, 31):
        ws2.cell(row=i, column=1, value=i)
    ws["C4"] = "=SUM(数据!A1:A30)"  # 高扇出（30 > 默认阈值 20）
    ws["C5"] = "=1+1"               # 孤立公式
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_audit_endpoint():
    r = client.post("/api/upload", files={"file": ("a.xlsx", make_audit_xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    wid = r.json()["workbook_id"]

    r = client.get(f"/api/workbooks/{wid}/audit")
    assert r.status_code == 200
    d = r.json()
    assert d["workbook_id"] == wid
    codes = set(d["counts_by_code"])
    assert codes == {"missing_sheet", "empty_ref", "out_of_range",
                     "high_fanout", "orphan_formula"}
    assert d["fanout_threshold"] == 20
    # 每条问题的必备字段
    for it in d["issues"]:
        assert {"code", "level", "sheet", "cell", "cell_display",
                "message", "evidence"} <= set(it)
    # 错误级排在最前
    assert d["issues"][0]["level"] == "error"


def test_audit_filters_and_threshold():
    r = client.post("/api/upload", files={"file": ("a.xlsx", make_audit_xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    wid = r.json()["workbook_id"]

    r = client.get(f"/api/workbooks/{wid}/audit", params={"level": "error"})
    assert {i["code"] for i in r.json()["issues"]} == {"missing_sheet"}

    r = client.get(f"/api/workbooks/{wid}/audit",
                   params=[("level", "warning"), ("level", "info")])
    assert r.json()["filters"]["levels"] == ["warning", "info"]

    r = client.get(f"/api/workbooks/{wid}/audit", params={"sheet": "S1"})
    assert r.status_code == 200
    # empty_ref 归属目标单元格 S1!B1；缺失工作表类归属公式所在表 S1
    assert "missing_sheet" in r.json()["counts_by_code"]

    r = client.get(f"/api/workbooks/{wid}/audit",
                   params={"fanout_threshold": "100"})
    assert "high_fanout" not in r.json()["counts_by_code"]
    r = client.get(f"/api/workbooks/{wid}/audit",
                   params={"fanout_threshold": "10"})
    assert "high_fanout" in r.json()["counts_by_code"]


def test_audit_bad_params_are_4xx():
    r = client.post("/api/upload", files={"file": ("a.xlsx", make_audit_xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    wid = r.json()["workbook_id"]

    assert client.get(f"/api/workbooks/{wid}/audit",
                      params={"level": "fatal"}).status_code == 400
    assert client.get(f"/api/workbooks/{wid}/audit",
                      params={"fanout_threshold": "abc"}).status_code == 400
    assert client.get(f"/api/workbooks/{wid}/audit",
                      params={"fanout_threshold": "0"}).status_code == 400
    assert client.get(f"/api/workbooks/{wid}/audit",
                      params={"fanout_threshold": "10001"}).status_code == 400
    assert client.get(f"/api/workbooks/{wid}/audit",
                      params={"fanout_threshold": "-3"}).status_code == 400
    assert client.get(f"/api/workbooks/{wid}/audit",
                      params={"sheet": "不存在"}).status_code == 404


def test_audit_missing_workbook_is_404():
    assert client.get("/api/workbooks/deadbeefab/audit").status_code == 404
    assert client.get("/api/workbooks/deadbeefab/audit/export").status_code == 404


def test_audit_export():
    r = client.post("/api/upload", files={"file": ("a.xlsx", make_audit_xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    wid = r.json()["workbook_id"]
    r = client.get(f"/api/workbooks/{wid}/audit/export")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    assert "dependency_audit_" in r.headers["content-disposition"]
    d = r.json()
    assert d["workbook_id"] == wid and d["filename"] == "a.xlsx"
    assert "issues" in d["audit"] and d["audit"]["total"] > 0
    # 高扇出证据在导出中不截断
    fanout = next(i for i in d["audit"]["issues"] if i["code"] == "high_fanout")
    assert len(fanout["evidence"]) == 30


def test_full_export_includes_audit():
    r = client.post("/api/upload", files={"file": ("a.xlsx", make_audit_xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    wid = r.json()["workbook_id"]
    d = client.get(f"/api/workbooks/{wid}/export").json()
    assert "audit" in d and "issues" in d["audit"]


def make_names_xlsx() -> bytes:
    from openpyxl.workbook.defined_name import DefinedName
    wb = Workbook()
    p = wb.active
    p.title = "参数"
    p["B2"] = 9
    c = wb.create_sheet("计算")
    c["B1"] = "=折扣率*2"          # 工作簿级名称 -> 参数!B2
    c["B2"] = "=幽灵名称"          # 未定义名称
    wb.defined_names["折扣率"] = DefinedName("折扣率", attr_text="参数!$B$2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_upload_reports_defined_name_counts():
    r = client.post("/api/upload", files={"file": ("n.xlsx", make_names_xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    d = r.json()
    assert d["defined_name_count"] == 1
    assert d["workbook_scoped_count"] == 1 and d["sheet_scoped_count"] == 0
    assert d["name_usage_count"] == 2


def test_cells_searchable_by_name_and_definition():
    wid = client.post("/api/upload", files={"file": ("n.xlsx", make_names_xlsx(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                      ).json()["workbook_id"]
    r = client.get(f"/api/workbooks/{wid}/cells", params={"q": "折扣率"})
    cells = {c["cell"] for c in r.json()["cells"]}
    assert "计算!B1" in cells
    r = client.get(f"/api/workbooks/{wid}/cells", params={"q": "$B$2"})
    assert "计算!B1" in {c["cell"] for c in r.json()["cells"]}
    # 单元格行携带名称留痕
    row = next(c for c in client.get(f"/api/workbooks/{wid}/cells",
                                     params={"q": "折扣率"}).json()["cells"]
               if c["cell"] == "计算!B1")
    assert row["names"][0]["name"] == "折扣率"
    assert row["names"][0]["status"] == "resolved"


def test_impact_through_name_and_root_names_field():
    wid = client.post("/api/upload", files={"file": ("n.xlsx", make_names_xlsx(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                      ).json()["workbook_id"]
    # 下游：参数!B2 经名称影响 计算!B1
    r = client.get(f"/api/workbooks/{wid}/impact", params={"cell": "参数!B2"})
    assert "计算!B1" in {i["cell"] for i in r.json()["impacted"]}
    # 上游 + 根名称字段
    r = client.get(f"/api/workbooks/{wid}/impact",
                   params={"cell": "计算!B1", "direction": "upstream"})
    d = r.json()
    assert "参数!B2" in {i["cell"] for i in d["impacted"]}
    assert d["root_names"][0]["name"] == "折扣率"
    assert d["root_names"][0]["definition"] == "参数!$B$2"


def test_audit_export_includes_name_inventory_and_usages():
    wid = client.post("/api/upload", files={"file": ("n.xlsx", make_names_xlsx(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                      ).json()["workbook_id"]
    d = client.get(f"/api/workbooks/{wid}/audit/export").json()
    names = {n["name"]: n for n in d["defined_names"]}
    assert "折扣率" in names and names["折扣率"]["definition"] == "参数!$B$2"
    statuses = {u["ref_text"]: u["status"] for u in d["name_usages"]}
    assert statuses["折扣率"] == "resolved"
    assert statuses["幽灵名称"] == "undefined"
    # 完整导出同样保留名称
    ex = client.get(f"/api/workbooks/{wid}/export").json()
    assert "defined_names" in ex and "name_usages" in ex
    assert ex["formulas"]["计算!B2"]["names"][0]["status"] == "undefined"


def test_audit_undefined_name_endpoint():
    wid = client.post("/api/upload", files={"file": ("n.xlsx", make_names_xlsx(),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                      ).json()["workbook_id"]
    d = client.get(f"/api/workbooks/{wid}/audit").json()
    und = next(i for i in d["issues"] if i["code"] == "undefined_name")
    assert und["level"] == "error" and und["cell"] == "计算!B2"
    assert und["evidence"][0]["to_display"] == "幽灵名称"
    # error 级别筛选只含名称错误（该工作簿无缺失工作表）
    errs = client.get(f"/api/workbooks/{wid}/audit",
                      params={"level": "error"}).json()["issues"]
    assert {i["code"] for i in errs} == {"undefined_name"}
